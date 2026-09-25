from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "presentacion_justificacion_zuly_gonzalez.pptx"

slides = [
    (
        "Justificacion del Proyecto",
        [
            "Aportes implementados por Zuly Gonzalez",
            "Proyecto: representacion lexical de tweets en espanol",
            "Enfoque: cumplimiento de criterios y entregables",
        ],
    ),
    (
        "Objetivo",
        [
            "Implementar y validar tres caracteristicas lexicas nuevas:",
            "- Dudas (doubt_count)",
            "- Regionalismos (regionalism_count)",
            "- Intensificadores (intensifier_count)",
        ],
    ),
    (
        "Alcance",
        [
            "Incluye solo cambios reales de esta iteracion.",
            "Se excluye la linea base previa del repositorio.",
            "Soporte en espanol unicamente.",
        ],
    ),
    (
        "Archivos clave",
        [
            "logic/feature_extraction.py",
            "logic/lexical_features.py",
            "tests/test_nuevas_features.py",
            "scripts/procesar_tweets_features.py",
            "arranque_entregable.ps1",
        ],
    ),
    (
        "Implementacion",
        [
            "Vector de salida: 8 caracteristicas.",
            "Nuevas caracteristicas integradas:",
            "- doubt_count",
            "- regionalism_count",
            "- intensifier_count",
        ],
    ),
    (
        "Metodos creados",
        [
            "get_features_nuevas(message)",
            "doubt_count(message, lexical)",
            "lexical_word_count(message, lexical_words)",
            "_normalize_text_for_match(text)",
            "process_split(path_in, path_out, fe)",
        ],
    ),
    (
        "Validacion",
        [
            "Pruebas de nuevas features en espanol: 16 passed.",
            "Arranque completo sin errores con arranque_entregable.ps1.",
            "Procesamiento TASS train/test completado.",
        ],
    ),
    (
        "Resultados",
        [
            "CSV generados:",
            "- data/tass/tass2018_es_train_features.csv",
            "- data/tass/tass2018_es_test_features.csv",
            "Promedios train: doubt_count 0.0813, regionalism_count 0.0466, intensifier_count 0.1647",
        ],
    ),
    (
        "Cumplimiento de criterios",
        [
            "Correctitud funcional: pruebas en verde.",
            "Reproducibilidad: script unico de arranque.",
            "Trazabilidad: metodos y archivos identificados.",
            "Evidencia objetiva: salidas CSV y metricas.",
        ],
    ),
    (
        "Cierre",
        [
            "Las tres caracteristicas fueron implementadas y validadas.",
            "El flujo completo quedo operativo.",
            "La justificacion tecnica queda soportada por evidencia ejecutada.",
        ],
    ),
]


prs = Presentation()

for title, bullets in slides:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print(f"Presentacion generada: {OUT}")
